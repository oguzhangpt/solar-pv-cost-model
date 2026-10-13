"""Independent reference calculations with LibreOffice.

Excel's cached results only cover the configuration the workbook was saved in
(China, all stages domestic). To validate the application's engine for other
countries, sourcing options and edited inputs, this script writes temporary
*copies* of the workbook with the configuration applied, lets LibreOffice
recalculate them headlessly, and stores the recalculated outputs as JSON
fixtures in validation/fixtures/. The original workbook is never modified.

Usage:  python3 tools/libreoffice_recalc.py
Requires: LibreOffice (soffice) on PATH, openpyxl.
"""
import json
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parents[1]
WB = ROOT / "IRENA_TSSC_Solar_PV_Manufacturing_Cost_Model_v1.xlsx"
OUT = ROOT / "validation" / "fixtures" / "libreoffice_recalc.json"

CALC = "Cost Model_Calculations"
YEARS = "DEFGHIJKLMN"
# Outputs captured for every configuration: stage prices (with overheads),
# totals and module power, for all years and both technologies.
ROWS = [14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 38, 39, 40, 41, 42, 43, 44, 45, 46, 47, 48,
        59, 60, 61, 62, 63, 64, 65, 66, 67, 68, 69, 77, 81, 82, 83, 84, 85, 86, 87, 88, 89, 90, 91, 92, 95, 96,
        110, 111, 112, 113, 114, 115, 116, 117, 118, 119, 120, 121, 122, 134, 135, 136, 137, 138, 139, 140,
        141, 142, 143, 144, 155, 156, 157, 158, 159, 160, 161, 162, 163, 164, 165, 173,
        177, 178, 179, 180, 181, 182, 183, 184, 185, 186, 187, 188, 191, 192]
OUTPUT_KEYS = [f"{CALC}!{c}{r}" for r in ROWS for c in YEARS] + \
    [f"Output_Visualization!C{r}" for r in list(range(199, 204)) + list(range(207, 212)) +
     list(range(215, 220)) + list(range(223, 228))] + \
    [f"Inputs!D{r}" for r in list(range(11, 17)) + list(range(23, 33)) + list(range(41, 47)) + list(range(53, 63))]

SOURCE_CELLS = {"topcon": ("C8", "C19", "C20"), "mono": ("C38", "C49", "C50")}
OPTIONS = {"Germany": ["Domestic", "China"], "India": ["Domestic", "China"], "Australia": ["Domestic", "China"],
           "Vietnam": ["Domestic", "China"], "US": ["Domestic", "Vietnam"], "China": ["Domestic"]}


def configurations():
    cfgs = []
    for country, opts in OPTIONS.items():
        imp = opts[-1]
        cfgs.append({"id": f"{country}-domestic", "country": country,
                     "sources": {"topcon": ["Domestic"] * 3, "mono": ["Domestic"] * 3}, "inputs": {}})
        if imp != "Domestic":
            cfgs.append({"id": f"{country}-all-imported", "country": country,
                         "sources": {"topcon": [imp] * 3, "mono": [imp] * 3}, "inputs": {}})
            cfgs.append({"id": f"{country}-mixed", "country": country,
                         "sources": {"topcon": ["Domestic", imp, "Domestic"], "mono": [imp, "Domestic", imp]},
                         "inputs": {}})
    # Single-input change (other inputs constant): +20% electricity price, China.
    cfgs.append({"id": "China-elec+20pct", "country": "China",
                 "sources": {"topcon": ["Domestic"] * 3, "mono": ["Domestic"] * 3},
                 "inputs": {"E13": 0.065 * 1.2, "E28": 0.08 * 1.2, "E43": 0.08 * 1.2, "E58": 0.08 * 1.2}})
    # Capacity change (drives ESG certification cost per Wp).
    cfgs.append({"id": "Germany-capacity-override", "country": "Germany",
                 "sources": {"topcon": ["Domestic"] * 3, "mono": ["Domestic"] * 3},
                 "inputs": {"E23": 1e9, "E24": 3e9, "E53": 1e9, "E54": 3e9}})
    return cfgs


def main():
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) not found")
    results = {}
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        jobs = []
        for cfg in configurations():
            wb = openpyxl.load_workbook(WB)
            ws = wb["Inputs"]
            ws["C3"] = cfg["country"]
            for tech, cells in SOURCE_CELLS.items():
                for cell, val in zip(cells, cfg["sources"][tech]):
                    ws[cell] = val
            for cell, val in cfg["inputs"].items():
                ws[cell] = val
            src = tmp / f"{cfg['id']}.xlsx"
            wb.save(src)  # openpyxl drops cached values -> LibreOffice must recalculate
            jobs.append((cfg, src))
        outdir = tmp / "out"
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir)]
                       + [str(s) for _, s in jobs], check=True, capture_output=True, timeout=600)
        for cfg, src in jobs:
            wv = openpyxl.load_workbook(outdir / src.name, data_only=True)
            vals = {}
            for key in OUTPUT_KEYS:
                sheet, a = key.rsplit("!", 1)
                v = wv[sheet][a].value
                vals[key] = v if not isinstance(v, str) else v
            results[cfg["id"]] = {"config": cfg, "values": vals}
            print(f"recalculated {cfg['id']}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    version = subprocess.run([soffice, "--version"], capture_output=True, text=True).stdout.strip()
    OUT.write_text(json.dumps({"libreoffice": version, "configs": results}, indent=1))
    print(f"Wrote {OUT.relative_to(ROOT)} ({len(results)} configurations)")


if __name__ == "__main__":
    main()
