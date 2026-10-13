"""Workbook ingestion layer.

Reads the original IRENA workbook (read-only, never modified) and exports every
populated cell -- constant value, formula text and Excel's cached result -- plus
structural metadata (sheets, visibility, protection, defined names, data
validation lists, charts) into ``app/data/model_data.js``.

The browser application and the Node validation suite both execute the exported
formulas with the same formula engine (app/assets/js/formula.js), so the
application reproduces the workbook's own calculation logic rather than a
hand-written approximation.

Usage:  python3 tools/extract_workbook.py [path/to/workbook.xlsx]
"""
import datetime
import hashlib
import json
import re
import sys
import warnings
import zipfile
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")  # openpyxl warns about the x14 validation extension

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WB = ROOT / "IRENA_TSSC_Solar_PV_Manufacturing_Cost_Model_v1.xlsx"
OUT = ROOT / "app" / "data" / "model_data.js"


def formula_text(value):
    """Return formula text for plain and array formulas, else None."""
    if hasattr(value, "text"):  # openpyxl ArrayFormula
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def jsonable(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    if isinstance(v, float) and v != v:  # NaN
        return None
    return v


def extract(path: Path):
    raw = path.read_bytes()
    wb_f = openpyxl.load_workbook(path)                  # formulas
    wb_v = openpyxl.load_workbook(path, data_only=True)  # cached results

    zf = zipfile.ZipFile(path)
    workbook_xml = zf.read("xl/workbook.xml").decode("utf8")

    sheets, cells = [], {}
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        n_vals = n_forms = 0
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                n_vals += 1
                key = f"{ws.title}!{c.coordinate}"
                f = formula_text(c.value)
                cached = jsonable(wv[c.coordinate].value)
                if f is not None:
                    n_forms += 1
                    cells[key] = {"f": f, "v": cached}
                    if hasattr(c.value, "text"):
                        cells[key]["array"] = True
                else:
                    cells[key] = {"v": jsonable(c.value)}
        hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
        hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
        sheets.append({
            "name": ws.title,
            "state": ws.sheet_state,
            "dimensions": ws.dimensions,
            "populated_cells": n_vals,
            "formula_cells": n_forms,
            "charts": len(ws._charts),
            "hidden_rows": hidden_rows,
            "hidden_cols": hidden_cols,
            "protected": bool(ws.protection.sheet),
        })

    # Data validation: openpyxl keeps the classic list validations; the x14
    # extension (country selector) is read straight from the sheet XML.
    validations = []
    for ws in wb_f.worksheets:
        for dv in ws.data_validations.dataValidation:
            validations.append({"sheet": ws.title, "sqref": str(dv.sqref),
                                "type": dv.type, "formula": dv.formula1})
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf8")
    for m in re.finditer(r'<sheet name="([^"]+)" sheetId="\d+"(?: state="\w+")? r:id="(rId\d+)"', workbook_xml):
        name, rid = m.groups()
        target = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels) or re.search(rf'Target="([^"]+)"[^>]*Id="{rid}"', rels)
        xml = zf.read("xl/" + target.group(1).lstrip("/").replace("xl/", "")).decode("utf8")
        for x in re.finditer(r"<x14:dataValidation [^>]*type=\"(\w+)\".*?<xm:f>([^<]+)</xm:f>.*?<xm:sqref>([^<]+)</xm:sqref>", xml, re.S):
            validations.append({"sheet": name, "sqref": x.group(3), "type": x.group(1),
                                "formula": x.group(2), "extension": "x14"})

    defined_names = {n: d.attr_text for n, d in wb_f.defined_names.items()}

    charts = []
    for name in sorted(zf.namelist()):
        if re.match(r"xl/charts/chart(Ex)?\d+\.xml$", name):
            xml = zf.read(name).decode("utf8")
            title = " ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)[:1])
            kind = "waterfall" if 'layoutId="waterfall"' in xml else (
                re.search(r"<c:(barChart|lineChart|pieChart|areaChart)", xml).group(1)
                if re.search(r"<c:(barChart|lineChart|pieChart|areaChart)", xml) else "unknown")
            refs = sorted(set(re.findall(r"<c:f>([^<]*)</c:f>", xml) + re.findall(r"<cx:f[^>]*>([^<]*)</cx:f>", xml)))
            charts.append({"part": name, "type": kind, "title": title, "refs": refs})

    calc = re.search(r"<calcPr([^/]*)/>", workbook_xml)
    return {
        "meta": {
            "file": path.name,
            "sha256": hashlib.sha256(raw).hexdigest(),
            "bytes": len(raw),
            "extracted_utc": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "calcPr": calc.group(1).strip() if calc else None,
            "workbook_protected": "<workbookProtection" in workbook_xml,
            "sheets": sheets,
        },
        "cells": cells,
        "definedNames": defined_names,
        "validations": validations,
        "charts": charts,
    }


def main():
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WB
    data = extract(path)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    OUT.write_text("/* Auto-generated by tools/extract_workbook.py from "
                   f"{data['meta']['file']} -- do not edit by hand. */\n"
                   f"window.PV_WORKBOOK = {payload};\n", encoding="utf8")
    print(f"Wrote {OUT.relative_to(ROOT)}: {len(data['cells'])} cells, "
          f"{sum(1 for c in data['cells'].values() if 'f' in c)} formulas")


if __name__ == "__main__":
    main()
